from __future__ import annotations

from pathlib import Path

from jean_finder.core.config import AppConfig
from jean_finder.core.database import Database
from jean_finder.core.indexer import Indexer
from jean_finder.core.searcher import Searcher


NEEDLE = "고유키워드ABCXYZ"


def _make_docx(path: Path, text: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def _make_xlsx(path: Path, text: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([text])
    wb.save(str(path))


def test_index_then_search(tmp_path: Path):
    docs = tmp_path / "docs"
    docs.mkdir()
    _make_docx(docs / "a.docx", f"앞에 단어가 있고 {NEEDLE} 뒤에도 단어가 있다")
    _make_xlsx(docs / "b.xlsx", f"엑셀파일에도 {NEEDLE} 들어있음")
    _make_docx(docs / "c.docx", "관계없는 텍스트")

    cfg = AppConfig(data_dir=tmp_path / "data")
    cfg.ensure_dirs()
    db = Database(cfg.db_path)
    try:
        indexer = Indexer(db)
        progress = indexer.index_roots([docs])
        assert progress.indexed == 3
        assert progress.failed == 0

        searcher = Searcher(db, snippet_radius=cfg.snippet_radius)
        hits = searcher.search(NEEDLE)
        assert len(hits) == 2
        for h in hits:
            assert NEEDLE in h.snippet

        # 미존재 검색어
        assert searcher.search("절대없는단어ZZZZZ") == []

        # 재스캔: mtime 동일 → 모두 skip
        progress2 = indexer.index_roots([docs])
        assert progress2.indexed == 0
        assert progress2.skipped == 3

        # delete_missing
        (docs / "c.docx").unlink()
        # 다시 스캔하지 않고 missing 정리만
        present = {str(p) for p in docs.iterdir()}
        removed = db.delete_missing(present)
        assert removed == 1
        assert db.stats()["total"] == 2
    finally:
        db.close()


def test_snippet_radius(tmp_path: Path):
    docs = tmp_path / "docs"
    docs.mkdir()
    long_text = ("X" * 200) + NEEDLE + ("Y" * 200)
    _make_docx(docs / "long.docx", long_text)

    cfg = AppConfig(data_dir=tmp_path / "data")
    db = Database(cfg.db_path)
    try:
        Indexer(db).index_roots([docs])
        hits = Searcher(db, snippet_radius=10).search(NEEDLE)
        assert len(hits) == 1
        snip = hits[0].snippet
        assert NEEDLE in snip
        # 앞뒤 10자 + needle, 양끝 ellipsis 가능
        cleaned = snip.replace("...", "")
        before, _, after = cleaned.partition(NEEDLE)
        assert len(before) <= 10
        assert len(after) <= 10
    finally:
        db.close()
